#!/usr/bin/env python3
"""Build the standard Industry Analyst Modeler workbook from a JSON contract."""

from __future__ import annotations

import argparse
import json
import math
from datetime import date
from pathlib import Path

try:
    from openpyxl import Workbook
    from openpyxl.comments import Comment
    from openpyxl.worksheet.datavalidation import DataValidation
    from openpyxl.formatting.rule import CellIsRule
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter
except ImportError as exc:  # pragma: no cover
    raise SystemExit("openpyxl is required: python -m pip install openpyxl") from exc


BLUE = "1F4E78"
LIGHT_BLUE = "D9EAF7"
LIGHT_GRAY = "E7E6E6"
LINK_BLUE = "0563C1"
GREEN = "C6E0B4"
YELLOW = "FFE699"
RED = "F4CCCC"


def is_formula(value: object) -> bool:
    return isinstance(value, str) and value.startswith("=")


def is_number(value: object) -> bool:
    return isinstance(value, (int, float)) and not isinstance(value, bool)


def is_missing(value: object) -> bool:
    return value is None or value == ""


def subject_maps(data: dict) -> tuple[dict[str, str], dict[str, str]]:
    return (
        {item["id"]: item["name"] for item in data["objects"]},
        {item["id"]: item["name"] for item in data["anchors"]},
    )


def rating_formula(score_cell: str, rules: list[dict]) -> str:
    if not rules:
        return '=""'
    ordered = sorted(rules, key=lambda item: float(item["min_score"]), reverse=True)
    formula = f'"{ordered[-1]["label"]}"'
    for rule in reversed(ordered[:-1]):
        formula = f'IF({score_cell}>={float(rule["min_score"]):.10g},"{rule["label"]}",{formula})'
    return "=" + formula


def dimension_groups(data: dict) -> list[tuple[dict, list[dict], float]]:
    groups = []
    for dimension in data["dimensions"]:
        indicators = [item for item in data["indicators"] if item["dimension_id"] == dimension["id"]]
        groups.append((dimension, indicators, sum(float(item["weight"]) for item in indicators)))
    return groups


def indicator_weight_refs(data: dict, members: list[dict] | None = None) -> list[str]:
    return [f"'指标详表'!${get_column_letter(data['indicators'].index(item) + 2)}$3"
            for item in (members if members is not None else data["indicators"])]


def dimension_weight_formula(data: dict, members: list[dict]) -> str:
    return "=SUM(" + ",".join(indicator_weight_refs(data, members)) + ")"


def weighted_formula(refs: list[str], weights: list[float | str], scale: float = 1) -> str:
    terms = [f"{weight:.10g}" if is_number(weight) else weight for weight in weights]
    numerator = "+".join(f"{weight}*{ref}" for ref, weight in zip(refs, terms))
    denominator = "+".join(terms)
    return f'=IF(AND(COUNT({",".join(refs)})={len(refs)},({denominator})>0),({numerator})/({denominator})*{scale:.10g},"")'


def score_note(data: dict, subject_type: str, subject: dict, indicator_id: str) -> str:
    sources = [source for source in data["sources"]
               if (source["subject_type"], source["subject_id"], source["indicator_id"]) ==
                  (subject_type, subject["id"], indicator_id)]
    reason = (subject.get("missing_reasons", {}).get(indicator_id, "") if is_missing(subject["scores"][indicator_id])
              else subject.get("score_reasons", {}).get(indicator_id, ""))
    return reason + "\n" + "\n".join(f"{source['id']}：{source['claim']}" for source in sources)


def editable(cell, note: str = "") -> None:
    cell.font = Font(color=LINK_BLUE)
    cell.fill = PatternFill("solid", fgColor=LIGHT_BLUE)
    if note:
        cell.comment = Comment(note, "评分依据")


def difference_formula(left: str, right: str, left_scale: int = 1) -> str:
    return f'=IF(COUNT({left},{right})=2,{left}*{left_scale}-{right},"")'


def coverage_label(data: dict, subject_type: str, subject_id: str) -> str:
    indicators = data["indicators"]
    source_by_pair: dict[tuple[str, str, str], list[dict]] = {}
    for source in data["sources"]:
        key = (source["subject_type"], source["subject_id"], source["indicator_id"])
        source_by_pair.setdefault(key, []).append(source)
    strong = estimate = gap = 0
    for indicator in indicators:
        states = {item.get("status", "") for item in source_by_pair[(subject_type, subject_id, indicator["id"])]}
        if states & {"核实数据", "二手数据"}:
            strong += 1
        elif "估算数据" in states:
            estimate += 1
        else:
            gap += 1
    coverage = f"核实/二手 {strong}/{len(indicators)}；估算 {estimate}；缺口/待确认 {gap}"
    if subject_type == "对象":
        obj = next(item for item in data["objects"] if item["id"] == subject_id)
        return coverage + "；" + obj["confidence"]
    return coverage


def load_contract(path: Path) -> dict:
    data = json.loads(path.read_text(encoding="utf-8-sig"))
    required = ("model", "dimensions", "indicators", "objects", "anchors", "sources", "assumptions")
    missing = [key for key in required if key not in data]
    if missing:
        raise ValueError(f"Missing top-level fields: {', '.join(missing)}")
    if not data["indicators"] or not data["objects"]:
        raise ValueError("At least one indicator and one object are required")
    if data["model"].get("score_scale") != 100:
        raise ValueError("Set model.score_scale=100 and convert all scores/rules/formulas to 0-100; keep raw-data units unchanged")
    weights = [float(item["weight"]) for item in data["indicators"]]
    if any(not math.isfinite(weight) or weight <= 0 for weight in weights):
        raise ValueError("Indicator weights must be finite and positive")
    if abs(sum(weights) - 1.0) > 1e-8:
        raise ValueError(f"Indicator weights must total 1.0, got {sum(weights):.8f}")
    ids = [item["id"] for item in data["indicators"]]
    if len(ids) != len(set(ids)):
        raise ValueError("Indicator ids must be unique")
    dimension_ids = [item["id"] for item in data["dimensions"]]
    if not dimension_ids or len(dimension_ids) != len(set(dimension_ids)):
        raise ValueError("Dimensions need unique ids")
    for dimension in data["dimensions"]:
        if not all(str(dimension.get(field, "")).strip() for field in ("name", "question", "reason")):
            raise ValueError("Each dimension needs name, question, and research-based reason")
    for indicator in data["indicators"]:
        if indicator.get("dimension_id") not in dimension_ids:
            raise ValueError(f"Indicator {indicator['id']} needs a valid dimension_id")
        if not all(str(indicator.get(field, "")).strip() for field in ("name", "direction", "reason", "score_rule")):
            raise ValueError(f"Indicator {indicator['id']} needs name, direction, reason, and score_rule")
    if any(not indicators for _, indicators, _ in dimension_groups(data)):
        raise ValueError("Every dimension needs at least one indicator")
    object_ids = [item.get("id") for item in data["objects"]]
    anchor_ids = [item.get("id") for item in data["anchors"]]
    if len(object_ids) != len(set(object_ids)) or len(anchor_ids) != len(set(anchor_ids)):
        raise ValueError("Object and anchor ids must each be unique")
    if not data["anchors"]:
        raise ValueError("At least one comparison anchor is required")
    for obj in data["objects"]:
        if not str(obj.get("confidence", "")).strip():
            raise ValueError(f"Object {obj.get('name')} needs confidence assessment covering audit and conflicts")
    for collection in ("objects", "anchors"):
        for item in data[collection]:
            unknown = set(item.get("scores", {})) - set(ids)
            if unknown:
                raise ValueError(f"Unknown indicators in {item.get('name')}: {sorted(unknown)}")
            missing_scores = set(ids) - set(item.get("scores", {}))
            if missing_scores:
                raise ValueError(f"Missing scores in {item.get('name')}: {sorted(missing_scores)}")
            for indicator_id, score in item["scores"].items():
                if is_missing(score):
                    if item.get("status") not in {"暂不评分", "不可比较"} or not item.get("missing_reasons", {}).get(indicator_id):
                        raise ValueError(f"Missing score for {item['name']} / {indicator_id} needs status and missing_reasons")
                    continue
                if not is_formula(score) and (not is_number(score) or not 0 <= float(score) <= 100):
                    raise ValueError(f"Score for {item.get('name')} / {indicator_id} must be 0-100 or an Excel formula")
                if not str(item.get("score_reasons", {}).get(indicator_id, "")).strip():
                    raise ValueError(f"Score for {item['name']} / {indicator_id} needs a score_reasons explanation")
            if collection == "objects" and any(is_missing(score) for score in item["scores"].values()) and item.get("status") not in {"暂不评分", "不可比较"}:
                raise ValueError(f"Incomplete object {item['name']} cannot be ranked")
            if collection == "objects" and all(not is_missing(score) for score in item["scores"].values()) and item.get("status") in {"暂不评分", "不可比较"}:
                raise ValueError(f"Object {item['name']} has complete scores but an unscored status")
    rules = data["model"].get("rating_rules", [])
    rule_thresholds = [float(item["min_score"]) for item in rules]
    if len(rule_thresholds) != len(set(rule_thresholds)) or any("label" not in item for item in rules):
        raise ValueError("rating_rules need unique min_score values and labels")
    source_ids = [item.get("id") for item in data["sources"]]
    if not all(source_ids) or len(source_ids) != len(set(source_ids)):
        raise ValueError("Source ids must be unique")
    valid_subjects = {("对象", item["id"]) for item in data["objects"]} | {("锚点", item["id"]) for item in data["anchors"]}
    source_pairs: dict[tuple[str, str, str], list[dict]] = {}
    for source in data["sources"]:
        pair = (source.get("subject_type"), source.get("subject_id"), source.get("indicator_id"))
        if pair[:2] not in valid_subjects or pair[2] not in ids:
            raise ValueError(f"Invalid source mapping: {pair}")
        if source.get("status") not in {"核实数据", "二手数据", "估算数据", "数据缺口", "待人工确认"}:
            raise ValueError(f"Source {source['id']} has an invalid data status")
        if not str(source.get("claim", "")).strip() or not (str(source.get("url", "")).strip() or str(source.get("notes", "")).strip()):
            raise ValueError(f"Source {source['id']} needs claim, status, and URL or source/assumption note")
        source_pairs.setdefault(pair, []).append(source)
    for subject_type, collection in (("对象", data["objects"]), ("锚点", data["anchors"])):
        for item in collection:
            for indicator_id in ids:
                pair = (subject_type, item["id"], indicator_id)
                if pair not in source_pairs:
                    raise ValueError(f"Missing source mapping for {subject_type} {item['id']} / {indicator_id}")
                score = item["scores"][indicator_id]
                if is_missing(score):
                    if not any(source.get("status") == "数据缺口" for source in source_pairs[pair]):
                        raise ValueError(f"Missing score for {item['name']} / {indicator_id} needs a 数据缺口 source")
                if data["model"].get("status") in {"正式模型", "formal"} and subject_type == "对象" and not is_missing(score):
                    if not any(source.get("status") in {"核实数据", "二手数据", "估算数据"} for source in source_pairs[pair]):
                        raise ValueError(f"Formal score for {item['name']} / {indicator_id} has only unconfirmed or gap sources")
    comparisons = data.get("anchor_comparisons", [])
    scored_ids = {item["id"] for item in data["objects"] if all(not is_missing(score) for score in item["scores"].values())}
    if not scored_ids.issubset({item.get("object_id") for item in comparisons}):
        raise ValueError("Every scored object needs an anchor comparison")
    for item in comparisons:
        if item.get("object_id") not in object_ids or item.get("anchor_id") not in anchor_ids or not str(item.get("note", "")).strip():
            raise ValueError(f"Invalid or unexplained anchor comparison: {item}")
    if data["model"].get("mode") == "组合加权":
        if len(data["objects"]) < 2:
            raise ValueError("组合加权 needs at least two child objects")
        if any(any(is_missing(score) for score in item["scores"].values()) for item in data["objects"]):
            raise ValueError("组合加权 needs complete child scores; use a bounded scenario for missing data")
        combination = data.get("combination", {})
        components = combination.get("components", [])
        if {item.get("object_id") for item in components} != set(object_ids) or len(components) != len(object_ids):
            raise ValueError("Combination components must match objects exactly")
        if abs(sum(float(item.get("weight", 0)) for item in components) - 1) > 1e-8:
            raise ValueError("Combination weights must total 1.0")
        if any(not (0 <= float(item["weight"]) <= 1) or not str(item.get("source", "")).strip() for item in components):
            raise ValueError("Each combination weight needs a valid value and market-size source")
    return data


def style_title(ws, text: str, end_col: int) -> None:
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=end_col)
    cell = ws.cell(1, 1, text)
    cell.font = Font(size=16, bold=True, color="FFFFFF")
    cell.fill = PatternFill("solid", fgColor=BLUE)
    cell.alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)
    ws.row_dimensions[1].height = 42


def style_header(row) -> None:
    for cell in row:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor=BLUE)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)


def set_widths(ws, widths: dict[int, float]) -> None:
    for column, width in widths.items():
        ws.column_dimensions[get_column_letter(column)].width = width


def score_formula(indicators: list[dict], row: int, start_col: int = 2) -> str:
    refs = [f"{get_column_letter(start_col + i)}{row}" for i in range(len(indicators))]
    weights = [f"${get_column_letter(start_col + i)}$3" for i in range(len(indicators))]
    return weighted_formula(refs, weights)


def build_detail(wb: Workbook, data: dict) -> tuple[dict[str, int], int]:
    ws = wb.create_sheet("指标详表")
    indicators = data["indicators"]
    last_col = len(indicators) + 2
    style_title(ws, "指标详表｜蓝色格调整权重和目标评分；评分理由见批注及模型假设", last_col)
    headers = ["对象"] + [item["name"] for item in indicators] + ["综合得分（百分制）"]
    ws.append(headers)
    style_header(ws[2])
    ws.append(["权重"] + [float(item["weight"]) for item in indicators] + ["=SUM(B3:%s3)" % get_column_letter(last_col - 1)])
    ws.append(["赋分规则"] + [item.get("score_rule", "") for item in indicators] + [""])
    object_rows: dict[str, int] = {}
    for obj in data["objects"]:
        row = ws.max_row + 1
        values = [obj["name"]] + [obj.get("scores", {}).get(item["id"]) for item in indicators]
        ws.append(values + [score_formula(indicators, row)])
        object_rows[obj["id"]] = row
        for col, item in enumerate(indicators, start=2):
            if not is_formula(ws.cell(row, col).value):
                editable(ws.cell(row, col), score_note(data, "对象", obj, item["id"]))
    for col, item in enumerate(indicators, start=2):
        editable(ws.cell(3, col), item["reason"])
        ws.cell(3, col).number_format = "0.0%"
    ws.cell(3, last_col).number_format = "0.0%"
    score_validation = DataValidation(type="decimal", operator="between", formula1=0, formula2=100, allow_blank=True)
    score_validation.showErrorMessage = True
    score_validation.error = "评分使用0—100分。"
    ws.add_data_validation(score_validation)
    score_validation.add(f"B5:{get_column_letter(last_col - 1)}{ws.max_row}")

    for cell in ws[4]:
        cell.fill = PatternFill("solid", fgColor=LIGHT_GRAY)
        cell.alignment = Alignment(wrap_text=True, vertical="top")
    ws.freeze_panes = "B5"
    ws.auto_filter.ref = f"A2:{get_column_letter(last_col)}{ws.max_row}"
    score_range = f"B5:{get_column_letter(last_col - 1)}{ws.max_row}"
    ws.conditional_formatting.add(score_range, CellIsRule(operator="greaterThanOrEqual", formula=["80"], fill=PatternFill("solid", fgColor=GREEN)))
    ws.conditional_formatting.add(score_range, CellIsRule(operator="between", formula=["50", "79.999999"], fill=PatternFill("solid", fgColor=YELLOW)))
    ws.conditional_formatting.add(score_range, CellIsRule(operator="lessThanOrEqual", formula=["49.999999"], fill=PatternFill("solid", fgColor=RED)))
    set_widths(ws, {1: 24, **{i: 16 for i in range(2, last_col)}, last_col: 22})
    return object_rows, last_col


def build_overview(
    wb: Workbook,
    data: dict,
    object_rows: dict[str, int],
    detail_score_col: int,
    anchor_score_cells: dict[str, str],
) -> None:
    ws = wb.active
    ws.title = "模型总览"
    indicators = data["indicators"]
    end_col = max(len(indicators) + 6, len(data["objects"]) + len(data["anchors"]) + 2)
    style_title(ws, data["model"].get("name", "行业评分模型"), end_col)
    metadata = [
        ("分析目标", data["model"].get("purpose", "")),
        ("建模模式", data["model"].get("mode", "")),
        ("比较宇宙", data["model"].get("comparison_universe", "")),
        ("模型状态", data["model"].get("status", "")),
        ("制作方/版本", f"{data['model'].get('author', '未注明')} / {data['model'].get('version', '')}"),
        ("制作日期", data["model"].get("date", date.today().isoformat())),
    ]
    for key, value in metadata:
        ws.append([key, value])
        ws.cell(ws.max_row, 1).font = Font(bold=True)
    ws.append(["所有评分满分100分；分差单位为分；蓝色格可调，其他位置由公式联动"])
    ws.append(["分维度比较（百分制）", "维度权重"] + [item["name"] for item in data["objects"] + data["anchors"]])
    style_header(ws[ws.max_row])
    for offset, (dimension, members, weight) in enumerate(dimension_groups(data)):
        anchor_row = len(indicators) + 8 + offset
        values = []
        for col in range(3, 3 + len(data["objects"]) + len(data["anchors"])):
            ref = f"'锚点参照与赋分意义'!{get_column_letter(col)}{anchor_row}"
            values.append(f'=IF({ref}="","",{ref})')
        ws.append([dimension["name"], dimension_weight_formula(data, members)] + values)
        ws.cell(ws.max_row, 2).number_format = "0%"
        for col in range(3, 3 + len(values)):
            ws.cell(ws.max_row, col).number_format = "0.0"
    totals = []
    for col in range(6, 6 + len(data["objects"]) + len(data["anchors"])):
        ref = f"'锚点参照与赋分意义'!{get_column_letter(col)}{len(indicators) + 5}"
        totals.append(f'=IF({ref}="","",{ref})')
    ws.append(["综合得分（百分制）", "=SUM('指标详表'!B3:" + get_column_letter(len(indicators) + 1) + "3)"] + totals)
    ws.cell(ws.max_row, 2).number_format = "0.0%"
    for cell in ws[ws.max_row]:
        cell.fill = PatternFill("solid", fgColor=LIGHT_BLUE)
        cell.font = Font(bold=True)
    ws.append([])
    ws.append(["排名", "对象"] + [item["name"] for item in indicators] + ["综合得分（百分制）", "评级", "模型置信度", "模型状态"])
    header_row = ws.max_row
    style_header(ws[header_row])
    detail_col = get_column_letter(detail_score_col)
    overview_score_col = get_column_letter(len(indicators) + 3)
    first_result_row = header_row + 1
    last_result_row = header_row + len(data["objects"])
    for obj in data["objects"]:
        row = object_rows[obj["id"]]
        overview_row = ws.max_row + 1
        score_cell = f"{overview_score_col}{overview_row}"
        indicator_cells = [f'=IF(\'指标详表\'!{get_column_letter(index)}{row}="","",\'指标详表\'!{get_column_letter(index)}{row})' for index in range(2, detail_score_col)]
        ws.append([
            f'=IF({score_cell}="","",RANK({score_cell},${overview_score_col}${first_result_row}:${overview_score_col}${last_result_row},0))',
            obj["name"],
            *indicator_cells,
            f'=IF(\'指标详表\'!{detail_col}{row}="","",\'指标详表\'!{detail_col}{row})',
            f'=IF({score_cell}="","",{rating_formula(score_cell, data["model"].get("rating_rules", []))[1:]})',
            coverage_label(data, "对象", obj["id"]),
            obj.get("status", data["model"].get("status", "")),
        ])
    rules = sorted(data["model"].get("rating_rules", []), key=lambda item: float(item["min_score"]), reverse=True)
    score_range = f"{overview_score_col}{first_result_row}:{overview_score_col}{last_result_row}"
    if rules:
        ws.conditional_formatting.add(score_range, CellIsRule(operator="greaterThanOrEqual", formula=[str(rules[0]["min_score"])], fill=PatternFill("solid", fgColor=GREEN)))
    if len(rules) > 1:
        ws.conditional_formatting.add(score_range, CellIsRule(operator="between", formula=[str(rules[1]["min_score"]), str(float(rules[0]["min_score"]) - 1e-6)], fill=PatternFill("solid", fgColor=YELLOW)))
    ws.append([])
    ws.append(["锚点定位", "对象", "参照锚点", "综合分差（百分制）", "定位说明"])
    style_header(ws[ws.max_row])
    for comparison in data.get("anchor_comparisons", []):
        object_id = comparison.get("object_id")
        anchor_id = comparison.get("anchor_id")
        if object_id not in object_rows or anchor_id not in anchor_score_cells:
            raise ValueError(f"Invalid anchor comparison: {comparison}")
        detail_row = object_rows[object_id]
        ws.append([
            "",
            next(item["name"] for item in data["objects"] if item["id"] == object_id),
            next(item["name"] for item in data["anchors"] if item["id"] == anchor_id),
            difference_formula(f"'指标详表'!{detail_col}{detail_row}", anchor_score_cells[anchor_id]),
            comparison.get("note", ""),
        ])
    ws.append([])
    ws.append(["研究结论（基准版；调参后复核）"])
    style_header(ws[ws.max_row])
    for finding in data["model"].get("key_findings", []):
        ws.append([finding])
        ws.merge_cells(start_row=ws.max_row, start_column=1, end_row=ws.max_row, end_column=end_col)
        ws.cell(ws.max_row, 1).alignment = Alignment(wrap_text=True, vertical="top")
        ws.row_dimensions[ws.max_row].height = 42
    ws.append([])
    ws.append(["使用顺序", "看什么 / 改哪里"])
    style_header(ws[ws.max_row])
    for name, purpose in [
        ("模型总览", "看维度、总分、差距与研究结论。"),
        ("锚点参照与赋分意义", "并排比较目标、理想模型和历史参照；蓝色格可调参照评分。"),
        ("指标详表", "第3行改权重、第5行起改目标评分；维度分、总分、差距自动更新。"),
        ("模型假设", "看维度选择、理想条件和逐项评分理由。"),
        ("数据溯源", "查原始事实、日期与来源链接。"),
    ]:
        ws.append([name, purpose])
        ws.cell(ws.max_row, 1).hyperlink = f"#'{name}'!A1"
        ws.cell(ws.max_row, 2).alignment = Alignment(wrap_text=True)
        ws.row_dimensions[ws.max_row].height = 48
    ws.append(["调参提示", "权重合计100%；修改后核对关键结论，保存为自己的方案。"])
    ws.cell(ws.max_row, 2).alignment = Alignment(wrap_text=True)
    ws.row_dimensions[ws.max_row].height = 48
    ws.freeze_panes = "C10"
    set_widths(ws, {1: 26, 2: 28, **{column: 18 for column in range(3, len(indicators) + 3)}, len(indicators) + 3: 24, len(indicators) + 4: 22, len(indicators) + 5: 28, len(indicators) + 6: 24})


def build_anchors(wb: Workbook, data: dict, object_rows: dict[str, int], detail_score_col: int) -> dict[str, str]:
    ws = wb.create_sheet("锚点参照与赋分意义")
    anchors = data["anchors"]
    indicators = data["indicators"]
    first_anchor_col = 6 + len(data["objects"])
    subjects = data["objects"] + anchors
    end_col = max(len(subjects) + 5, len(indicators) + 5)
    groups = dimension_groups(data)
    dimension_by_id = {dim["id"]: (dim, weight) for dim, _, weight in groups}
    style_title(ws, "锚点参照与赋分意义｜看目标与参照的逐项差距；蓝色格调整参照评分", end_col)
    ws.append(["比较宇宙", data["model"].get("comparison_universe", "")])
    ws.append(["所有评分满分100分；分差单位为分；蓝色格可调，其他位置由公式联动"])
    ws.append(["维度", "维度权重", "指标", "维度内权重", "全局权重"] + [item["name"] for item in subjects])
    style_header(ws[4])
    ws.row_dimensions[4].height = 36
    for index, indicator in enumerate(indicators, start=2):
        dimension, weight = dimension_by_id[indicator["dimension_id"]]
        values = []
        for obj in data["objects"]:
            ref = f"'指标详表'!{get_column_letter(index)}{object_rows[obj['id']]}"
            values.append(f'=IF({ref}="","",{ref})')
        row = ws.max_row + 1
        members = [item for item in indicators if item["dimension_id"] == indicator["dimension_id"]]
        ws.append([dimension["name"], dimension_weight_formula(data, members), indicator["name"],
                   f'=IF(B{row}>0,E{row}/B{row},"")', "=" + indicator_weight_refs(data, [indicator])[0]]
                  + values + [anchor["scores"][indicator["id"]] for anchor in anchors])
        for col, anchor in enumerate(anchors, start=first_anchor_col):
            if not is_formula(ws.cell(row, col).value):
                editable(ws.cell(row, col), score_note(data, "锚点", anchor, indicator["id"]))
        for col in (2, 4, 5):
            ws.cell(ws.max_row, col).number_format = "0%"
        for col in range(6, 6 + len(subjects)):
            ws.cell(ws.max_row, col).number_format = "0.0"
    anchor_range = f"F5:{get_column_letter(len(subjects) + 5)}{ws.max_row}"
    ws.conditional_formatting.add(anchor_range, CellIsRule(operator="greaterThanOrEqual", formula=["80"], fill=PatternFill("solid", fgColor=GREEN)))
    ws.conditional_formatting.add(anchor_range, CellIsRule(operator="between", formula=["50", "79.999999"], fill=PatternFill("solid", fgColor=YELLOW)))
    ws.conditional_formatting.add(anchor_range, CellIsRule(operator="lessThanOrEqual", formula=["49.999999"], fill=PatternFill("solid", fgColor=RED)))
    total_row = ws.max_row + 1
    ws.cell(total_row, 1, "综合得分（百分制）")
    anchor_score_cells: dict[str, str] = {}
    for col, subject in enumerate(subjects, start=6):
        refs = [f"{get_column_letter(col)}{row}" for row in range(5, total_row)]
        ws.cell(total_row, col, weighted_formula(refs, indicator_weight_refs(data)))
        ws.cell(total_row, col).number_format = "0.0"
        if col >= first_anchor_col:
            anchor_score_cells[subject["id"]] = f"'锚点参照与赋分意义'!{get_column_letter(col)}{total_row}"
    for cell in ws[total_row]:
        cell.font = Font(bold=True)
        cell.fill = PatternFill("solid", fgColor=LIGHT_BLUE)
    ws.append([])
    ws.append(["维度汇总（百分制）", "维度权重"] + [item["name"] for item in subjects])
    style_header(ws[ws.max_row])
    for dimension, members, weight in groups:
        row = ws.max_row + 1
        ws.cell(row, 1, dimension["name"])
        ws.cell(row, 2, dimension_weight_formula(data, members)).number_format = "0%"
        for col in range(3, 3 + len(subjects)):
            refs = [f"{get_column_letter(col + 3)}{5 + indicators.index(item)}" for item in members]
            ws.cell(row, col, weighted_formula(refs, indicator_weight_refs(data, members)))
            ws.cell(row, col).number_format = "0.0"
    ws.append([])
    ws.append(["分值含义", "说明"])
    style_header(ws[ws.max_row])
    for item in data["model"].get("scale_meanings", []):
        ws.append([item.get("range", ""), item.get("meaning", "")])
    ws.append([])
    ws.append(["锚点说明", "选择理由", "代表意义"])
    style_header(ws[ws.max_row])
    for anchor in anchors:
        ws.append([anchor["name"], anchor.get("reason", "待补"), anchor.get("meaning", "待补")])
    ws.append([])
    ws.append(["目标—锚点差距", "参照锚点"] + [f"{item['name']}差" for item in indicators] + ["综合分差（百分制）", "定位说明"])
    style_header(ws[ws.max_row])
    difference_header_row = ws.max_row
    object_names, _ = subject_maps(data)
    for comparison in data.get("anchor_comparisons", []):
        object_id = comparison.get("object_id")
        anchor_id = comparison.get("anchor_id")
        if object_id not in object_rows or anchor_id not in anchor_score_cells:
            continue
        detail_row = object_rows[object_id]
        anchor_index = next(index for index, anchor in enumerate(anchors, start=first_anchor_col) if anchor["id"] == anchor_id)
        differences = [difference_formula(f"'指标详表'!{get_column_letter(index)}{detail_row}", f"{get_column_letter(anchor_index)}{5 + offset}") for offset, index in enumerate(range(2, detail_score_col))]
        ws.append([object_names[object_id], next(item["name"] for item in anchors if item["id"] == anchor_id), *differences, difference_formula(f"'指标详表'!{get_column_letter(detail_score_col)}{detail_row}", f"{get_column_letter(anchor_index)}{total_row}"), comparison.get("note", "")])
    if ws.max_row > difference_header_row:
        difference_range = f"C{difference_header_row + 1}:{get_column_letter(len(indicators) + 2)}{ws.max_row}"
        ws.conditional_formatting.add(difference_range, CellIsRule(operator="lessThan", formula=["0"], fill=PatternFill("solid", fgColor=RED)))
        ws.conditional_formatting.add(difference_range, CellIsRule(operator="greaterThan", formula=["0"], fill=PatternFill("solid", fgColor=GREEN)))
    ws.append([])
    ws.append(["关键解读"])
    style_header(ws[ws.max_row])
    for finding in data["model"].get("key_findings", ["待补"]):
        ws.append([finding])
    set_widths(ws, {1: 28, **{i: 20 for i in range(2, end_col + 1)}, end_col: 28})
    ws.column_dimensions["C"].width = 32
    ws.freeze_panes = "F5"
    return anchor_score_cells


def build_sources(wb: Workbook, data: dict) -> None:
    ws = wb.create_sheet("数据溯源")
    headers = ["编号", "主体类型", "主体", "指标", "数据/主张", "数值", "单位", "时间", "数据状态", "审核评级", "URL", "备注"]
    style_title(ws, "数据溯源｜查原始事实和来源；评分理由见模型假设", len(headers))
    ws.append(headers)
    style_header(ws[2])
    object_names, anchor_names = subject_maps(data)
    indicator_names = {item["id"]: item["name"] for item in data["indicators"]}
    for source in data["sources"]:
        names = object_names if source["subject_type"] == "对象" else anchor_names
        ws.append([
            source.get("id", ""), source.get("subject_type", ""), names.get(source.get("subject_id"), ""), indicator_names.get(source.get("indicator_id"), ""),
            source.get("claim", ""), source.get("value", ""), source.get("unit", ""), source.get("date", ""), source.get("status", ""), source.get("rating", ""), source.get("url", ""), source.get("notes", ""),
        ])
        url = source.get("url", "")
        if url:
            cell = ws.cell(ws.max_row, 11)
            cell.hyperlink = url
            cell.font = Font(color=LINK_BLUE, underline="single")
    ws.freeze_panes = "A3"
    ws.auto_filter.ref = f"A2:L{ws.max_row}"
    set_widths(ws, {1: 12, 2: 12, 3: 20, 4: 20, 5: 38, 6: 14, 7: 12, 8: 14, 9: 20, 10: 14, 11: 42, 12: 34})


def build_assumptions(wb: Workbook, data: dict) -> None:
    ws = wb.create_sheet("模型假设")
    style_title(ws, "模型假设｜看维度选择、理想条件和逐项评分理由", 6)
    rows = [
        ("比较宇宙", data["model"].get("comparison_universe", "")),
        ("地域/时间/单位", data["model"].get("scope", "")),
        ("模型状态", data["model"].get("status", "")),
        ("版本", data["model"].get("version", "")),
        ("评分尺度", "0—100分；权重与评分均可调"),
        ("制作方/版本", f"{data['model'].get('author', '未注明')} / {data['model'].get('version', '')}"),
    ]
    for key, value in rows:
        ws.append([key, value])
        ws.cell(ws.max_row, 1).font = Font(bold=True)
    ws.append([])
    ws.append(["类别", "假设/参数", "取值", "理由", "确认状态"])
    style_header(ws[ws.max_row])
    for item in data["assumptions"]:
        ws.append([item.get("category", ""), item.get("item", ""), item.get("value", ""), item.get("reason", ""), item.get("status", "")])
    ws.append([])
    ws.append(["一级维度", "维度权重", "回答的问题", "研究依据与设置理由"])
    style_header(ws[ws.max_row])
    groups = dimension_groups(data)
    dimension_by_id = {dim["id"]: dim for dim, _, _ in groups}
    for dimension, members, weight in groups:
        ws.append([dimension["name"], dimension_weight_formula(data, members), dimension["question"], dimension["reason"]])
        ws.cell(ws.max_row, 2).number_format = "0%"
    ws.append([])
    ws.append(["指标", "全局权重", "方向", "赋分规则", "设置理由", "所属维度"])
    style_header(ws[ws.max_row])
    for item in data["indicators"]:
        ws.append([item["name"], "=" + indicator_weight_refs(data, [item])[0], item.get("direction", ""), item.get("score_rule", ""), item.get("reason", ""), dimension_by_id[item["dimension_id"]]["name"]])
    ws.append([])
    ws.append(["对象", "维度", "指标", "当前得分", "评分理由（基准版）", "事实与来源编号"])
    style_header(ws[ws.max_row])
    for subject_type, subjects in (("对象", data["objects"]), ("锚点", data["anchors"])):
        for subject in subjects:
            for index, indicator in enumerate(data["indicators"]):
                if subject_type == "对象":
                    ref = f"'指标详表'!{get_column_letter(index + 2)}{5 + data['objects'].index(subject)}"
                else:
                    col = 6 + len(data["objects"]) + data["anchors"].index(subject)
                    ref = f"'锚点参照与赋分意义'!{get_column_letter(col)}{5 + index}"
                sources = [source for source in data["sources"] if
                           (source["subject_type"], source["subject_id"], source["indicator_id"]) ==
                           (subject_type, subject["id"], indicator["id"])]
                facts = "\n".join(f"{source['id']}：{source['claim']}" for source in sources)
                reason = (subject.get("missing_reasons", {}).get(indicator["id"], "") if is_missing(subject["scores"][indicator["id"]])
                          else subject.get("score_reasons", {}).get(indicator["id"], ""))
                ws.append([subject["name"], dimension_by_id[indicator["dimension_id"]]["name"], indicator["name"],
                           f'=IF({ref}="","",{ref})', reason, facts])
                for cell in ws[ws.max_row]:
                    cell.alignment = Alignment(wrap_text=True, vertical="top")
                ws.row_dimensions[ws.max_row].height = 72
    ws.freeze_panes = "D2"
    set_widths(ws, {1: 22, 2: 22, 3: 28, 4: 40, 5: 52, 6: 60})


def combination_formula(letter: str, last: int) -> str:
    weights = f"$B$3:$B${last}"
    scores = f"{letter}3:{letter}{last}"
    count = last - 2
    return (f'=IF(AND(COUNT({weights})={count},MIN({weights})>=0,'
            f'SUM({weights})>0,COUNT({scores})={count}),'
            f'SUMPRODUCT({weights},{scores})/SUM({weights}),"")')


def build_combination(wb: Workbook, data: dict, object_rows: dict[str, int], detail_score_col: int) -> None:
    if data["model"].get("mode") != "组合加权":
        return
    ws = wb.create_sheet("组合加权")
    indicators = data["indicators"]
    end_col = len(indicators) + 4
    style_title(ws, "组合加权｜调子赛道占比，查看母赛道的整体表现", end_col)
    ws.append(["对象", "组合权重", "权重来源"] + [item["name"] for item in indicators] + ["综合得分（百分制）"])
    style_header(ws[2])
    names = {item["id"]: item["name"] for item in data["objects"]}
    for component in data["combination"]["components"]:
        object_id = component["object_id"]
        detail_row = object_rows[object_id]
        ws.append([names[object_id], float(component["weight"]), component["source"]] + [f"='指标详表'!{get_column_letter(col)}{detail_row}" for col in range(2, detail_score_col)] + [f"='指标详表'!{get_column_letter(detail_score_col)}{detail_row}"])
        editable(ws.cell(ws.max_row, 2), component["source"])
        ws.cell(ws.max_row, 2).number_format = "0.0%"
    last = ws.max_row
    ws.append(["整体（加权）", f"=SUM(B3:B{last})", data["combination"].get("weight_basis", "子品类市场规模占比") + "；按合计归一化；输入缺失或权重无效时结果留空"] + [combination_formula(get_column_letter(col), last) for col in range(4, end_col + 1)])
    for cell in ws[ws.max_row]:
        cell.fill = PatternFill("solid", fgColor=LIGHT_BLUE)
        cell.font = Font(bold=True)
    ws.freeze_panes = "D3"
    set_widths(ws, {1: 26, 2: 18, 3: 48, **{col: 20 for col in range(4, end_col + 1)}})


def build_audit(wb: Workbook, data: dict) -> None:
    audit = data.get("calculation_audit", [])
    if not audit:
        return
    ws = wb.create_sheet("计算审计")
    headers = ["对象", "指标", "子变量", "子变量权重", "原始数值", "数据来源编号", "标准化公式/证据阶梯", "子变量得分", "指标得分公式", "指标得分", "算法/标尺理由", "人工判断/参数", "审计备注"]
    style_title(ws, "计算审计｜复核原始值如何逐步换算成评分", len(headers))
    ws.append(headers)
    style_header(ws[2])
    keys = ("object", "indicator", "sub_variable", "sub_weight", "raw_value", "source_id", "normalization", "sub_score", "indicator_formula", "indicator_score", "reason", "manual_parameter", "notes")
    for item in audit:
        ws.append([item.get(key, "") for key in keys])
    ws.freeze_panes = "A3"
    ws.auto_filter.ref = f"A2:M{ws.max_row}"
    set_widths(ws, {i: 20 for i in range(1, len(headers) + 1)})


def main() -> int:
    parser = argparse.ArgumentParser()
    parser.add_argument("--input", required=True, type=Path)
    parser.add_argument("--output", required=True, type=Path)
    args = parser.parse_args()
    data = load_contract(args.input)
    wb = Workbook()
    object_rows, score_col = build_detail(wb, data)
    anchor_score_cells = build_anchors(wb, data, object_rows, score_col)
    build_overview(wb, data, object_rows, score_col, anchor_score_cells)
    build_sources(wb, data)
    build_assumptions(wb, data)
    build_combination(wb, data, object_rows, score_col)
    build_audit(wb, data)
    wb._sheets = [wb[name] for name in ["模型总览", "锚点参照与赋分意义", "指标详表", "模型假设", "数据溯源"] +
                  [name for name in ("组合加权", "计算审计") if name in wb.sheetnames]]
    wb.calculation.fullCalcOnLoad = True
    wb.calculation.forceFullCalc = True
    args.output.parent.mkdir(parents=True, exist_ok=True)
    wb.save(args.output)
    print(json.dumps({"output": str(args.output.resolve()), "sheets": wb.sheetnames, "objects": len(data["objects"]), "indicators": len(data["indicators"])}, ensure_ascii=False))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
