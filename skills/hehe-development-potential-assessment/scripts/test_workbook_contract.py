#!/usr/bin/env python3
"""Regression checks for workbook references, evidence, gaps, and combination."""

from __future__ import annotations

import copy
import json
import tempfile
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.utils import get_column_letter

from build_scoring_workbook import (
    build_anchors, build_assumptions, build_audit, build_combination,
    build_detail, build_overview, build_sources, load_contract,
)
from validate_scoring_workbook import find_row, validate


ROOT = Path(__file__).resolve().parents[1]


def main() -> None:
    base = json.loads((ROOT / "examples/model-input.example.json").read_text(encoding="utf-8-sig"))
    with tempfile.TemporaryDirectory() as temp:
        contract = Path(temp) / "input.json"
        workbook = Path(temp) / "model.xlsx"

        def build(data: dict) -> dict:
            contract.write_text(json.dumps(data, ensure_ascii=False), encoding="utf-8")
            checked = load_contract(contract)
            wb = Workbook()
            rows, col = build_detail(wb, checked)
            anchors = build_anchors(wb, checked, rows, col)
            build_overview(wb, checked, rows, col, anchors)
            build_sources(wb, checked)
            build_assumptions(wb, checked)
            build_combination(wb, checked, rows, col)
            build_audit(wb, checked)
            wb.save(workbook)
            return validate(workbook, contract)

        def corrupt(sheet: str, cell: str, value: object) -> None:
            wb = load_workbook(workbook)
            wb[sheet][cell] = value
            wb.save(workbook)
            result = validate(workbook, contract)
            assert not result["passed"], f"Corruption passed: {sheet}!{cell}"

        assert build(base)["passed"]
        wb = load_workbook(workbook)
        assert "$B$3" in wb["指标详表"]["E5"].value, "Total must use editable weights"
        assert all(source.get("unit") != "分" for source in base["sources"]), "Evidence uses facts, not copied scores"
        assert find_row(wb["模型假设"], "对象"), "Scoring reasons must be visible"
        rank_row = find_row(wb["模型总览"], "排名") + 1
        assert "RANK(" in wb["模型总览"].cell(rank_row, 1).value
        for ws in wb:
            for row in ws:
                for cell in row:
                    if isinstance(cell.value, str) and cell.value.startswith("="):
                        cell.value = cell.value.replace("'指标详表'!", "指标详表!").replace("'锚点参照与赋分意义'!", "锚点参照与赋分意义!")
        wb.save(workbook)
        assert validate(workbook, contract)["passed"], "Excel-style sheet reference normalization failed"
        assert build(base)["passed"]
        corrupt("模型总览", f"A{rank_row}", "=RANK.EQ(E10,$E$9:$E$10,0)")
        assert build(base)["passed"]
        wb = load_workbook(workbook)
        ws = wb["锚点参照与赋分意义"]
        row = next(cell.row for cells in ws for cell in cells if cell.value == "目标—锚点差距")
        corrupt(ws.title, f"C{row + 1}", "='指标详表'!B5-C4")
        assert build(base)["passed"]
        ws = load_workbook(workbook)["锚点参照与赋分意义"]
        total_row = next(cell.row for cells in ws for cell in cells if cell.value == "综合得分（百分制）")
        first_anchor = get_column_letter(6 + len(base["objects"]))
        corrupt(ws.title, f"{first_anchor}{total_row}", "=B5")
        assert build(base)["passed"]
        composite_col = get_column_letter(len(base["indicators"]) + 3)
        corrupt(ws.title, f"{composite_col}{row + 1}", "='指标详表'!D5-B5")
        assert build(base)["passed"]
        overview = load_workbook(workbook)["模型总览"]
        overview_row = next(cell.row for cells in overview for cell in cells if cell.value == "锚点定位")
        corrupt(overview.title, f"D{overview_row + 1}", "='指标详表'!D5-B5")
        assert build(base)["passed"]
        corrupt("数据溯源", "E3", None)
        assert build(base)["passed"]
        corrupt("数据溯源", "F3", 1)

        assert build(base)["passed"]
        ws = load_workbook(workbook)["锚点参照与赋分意义"]
        dimension_row = find_row(ws, "维度汇总（百分制）") + 1
        formula = ws.cell(dimension_row, 3).value
        assert "F5,F7" in formula and "F6" not in formula, "Nonadjacent dimension membership lost"
        corrupt(ws.title, f"C{dimension_row}", "=AVERAGE(F5:F7)*10")
        assert build(base)["passed"]
        corrupt(ws.title, "A5", "未分维度")
        assert build(base)["passed"]
        corrupt(ws.title, "D5", 0.4)

        missing = copy.deepcopy(base)
        missing["objects"][1]["scores"]["growth"] = None
        missing["objects"][1]["status"] = "不可比较"
        missing["objects"][1]["missing_reasons"] = {"growth": "缺少同期增长数据"}
        missing["sources"][2].update(value="", unit="", status="数据缺口", notes="缺少同期增长数据；待查行业协会资料")
        missing["anchor_comparisons"] = missing["anchor_comparisons"][:1]
        assert build(missing)["passed"]

        missing_anchor = copy.deepcopy(base)
        missing_anchor["model"].pop("rating_rules")
        missing_anchor["anchors"][1]["scores"]["growth"] = None
        missing_anchor["anchors"][1].update(status="暂不评分", missing_reasons={"growth": "同口径历史数据待补"})
        missing_anchor["sources"][6].update(value="", unit="", status="数据缺口", notes="同口径历史数据待补")
        assert build(missing_anchor)["passed"], "Partial anchors without rating rules must be supported"
        ws = load_workbook(workbook)["锚点参照与赋分意义"]
        diff_row = find_row(ws, "目标—锚点差距") + 1
        assert "COUNT(" in ws.cell(diff_row, 3).value, "Missing anchor must not become zero difference"

        for change in ("missing_dimensions", "unknown_dimension", "empty_dimension"):
            bad_dimension = copy.deepcopy(base)
            if change == "missing_dimensions":
                bad_dimension.pop("dimensions")
            elif change == "unknown_dimension":
                bad_dimension["indicators"][0]["dimension_id"] = "unknown"
            else:
                bad_dimension["dimensions"].append({"id": "empty", "name": "空维度", "question": "未实现", "reason": "测试"})
            contract.write_text(json.dumps(bad_dimension, ensure_ascii=False), encoding="utf-8")
            try:
                load_contract(contract)
            except ValueError:
                pass
            else:
                raise AssertionError(f"Invalid hierarchy accepted: {change}")

        combo = copy.deepcopy(base)
        combo["model"]["mode"] = "组合加权"
        combo["combination"] = {"components": [
            {"object_id": "segment_a", "weight": 625.3 / 769.4, "source": "示例市场规模口径A；待核实"},
            {"object_id": "segment_b", "weight": 144.1 / 769.4, "source": "示例市场规模口径B；待核实"},
        ]}
        assert build(combo)["passed"]
        combo["combination"]["weight_basis"] = "未来三年收入情景占比"
        assert build(combo)["passed"]
        sheet = load_workbook(workbook)["组合加权"]
        assert "未来三年收入情景占比" in sheet["C5"].value
        corrupt("组合加权", "D5", "=SUMPRODUCT($B$3:$B$4,D3:D4)")
        assert build(combo)["passed"]
        composite_col = get_column_letter(len(combo["indicators"]) + 4)
        corrupt("组合加权", f"{composite_col}3", "='指标详表'!D5")
        assert build(combo)["passed"]
        corrupt("组合加权", f"{composite_col}5", "=SUM(F3:F4)")

        bad = copy.deepcopy(base)
        bad["sources"][0].update(claim="", url="", notes="", value=1, status="核实数据")
        bad["model"]["status"] = "formal"
        contract.write_text(json.dumps(bad, ensure_ascii=False), encoding="utf-8")
        try:
            load_contract(contract)
        except ValueError:
            pass
        else:
            raise AssertionError("Unsupported formal score passed contract validation")

        for change in ("old_scale", "out_of_range", "missing_reason"):
            bad_input = copy.deepcopy(base)
            if change == "old_scale":
                bad_input["model"]["score_scale"] = 10
            elif change == "out_of_range":
                bad_input["objects"][0]["scores"]["growth"] = 101
            else:
                bad_input["objects"][0]["score_reasons"].pop("growth")
            contract.write_text(json.dumps(bad_input, ensure_ascii=False), encoding="utf-8")
            try:
                load_contract(contract)
            except ValueError:
                pass
            else:
                raise AssertionError(f"Invalid scoring input accepted: {change}")
        zero_score = copy.deepcopy(base)
        zero_score["objects"][0]["scores"]["risk"] = 0
        zero_score["objects"][0]["score_reasons"]["risk"] = "教学假设核心业务无法开展，对应0分。"
        zero_score["sources"][1]["claim"] = "教学假设核心业务无法开展。"
        assert build(zero_score)["passed"], "Zero is a valid score, distinct from missing data"

    print("PASS: 0-100 scale, editable weights, qualitative reasons, raw evidence, dimensions, ranking, gaps, optional ratings and combination")


if __name__ == "__main__":
    main()
