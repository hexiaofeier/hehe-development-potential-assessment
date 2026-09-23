#!/usr/bin/env python3
"""Validate structure and auditability of an Industry Analyst Modeler workbook."""

from __future__ import annotations

import argparse
import json
from pathlib import Path
from openpyxl.utils import get_column_letter

from build_scoring_workbook import (
    coverage_label, dimension_groups, difference_formula,
    load_contract as validate_input_contract, rating_formula, weighted_formula,
    indicator_weight_refs, dimension_weight_formula, score_formula,
)

try:
    from openpyxl import load_workbook
except ImportError as exc:  # pragma: no cover
    raise SystemExit("openpyxl is required: python -m pip install -r requirements.txt") from exc


REQUIRED_SHEETS = ["模型总览", "锚点参照与赋分意义", "指标详表", "数据溯源", "模型假设"]


def is_formula(value: object) -> bool:
    return isinstance(value, str) and value.startswith("=")


def normalized_formula(value: object) -> str:
    """Excel may remove optional sheet-name quotes when saving a workbook."""
    if not isinstance(value, str):
        return ""
    return value.replace("'指标详表'!", "指标详表!").replace("'锚点参照与赋分意义'!", "锚点参照与赋分意义!")


def find_row(ws, first_value: str) -> int | None:
    for row in ws.iter_rows():
        if row and row[0].value == first_value:
            return row[0].row
    return None


def find_header_row(ws, value: str) -> int | None:
    for row in ws.iter_rows():
        if any(cell.value == value for cell in row):
            return row[0].row
    return None


def header_map(ws, row: int) -> dict[str, int]:
    return {str(cell.value): cell.column for cell in ws[row] if cell.value not in (None, "")}


def load_contract(path: Path | None) -> dict | None:
    if path is None:
        return None
    return validate_input_contract(path)


def validate(path: Path, contract_path: Path | None = None) -> dict:
    wb = load_workbook(path, data_only=False)
    errors: list[str] = []
    warnings: list[str] = []
    try:
        contract = load_contract(contract_path)
    except (ValueError, KeyError, TypeError) as exc:
        contract = None
        errors.append(f"Invalid input contract: {exc}")
    missing = [name for name in REQUIRED_SHEETS if name not in wb.sheetnames]
    if missing:
        errors.append("Missing required sheets: " + ", ".join(missing))
        return {"passed": False, "errors": errors, "warnings": warnings, "sheets": wb.sheetnames}

    detail = wb["指标详表"]
    detail_headers = header_map(detail, 2)
    score_header = "综合得分（百分制）"
    if score_header not in detail_headers:
        errors.append(f"指标详表 lacks column: {score_header}")
    weight_row = find_row(detail, "权重")
    if not weight_row:
        errors.append("指标详表 lacks 权重 row")
    else:
        weights = [float(cell.value) for cell in detail[weight_row][1:-1] if isinstance(cell.value, (int, float))]
        if not weights or abs(sum(weights) - 1.0) > 1e-8:
            errors.append(f"Indicator weights must total 1.0; got {sum(weights):.8f}")
    rule_row = find_row(detail, "赋分规则")
    if not rule_row:
        errors.append("指标详表 lacks 赋分规则 row")
    else:
        empty_rules = [cell.column for cell in detail[rule_row][1:-1] if not cell.value]
        if empty_rules:
            warnings.append(f"Empty scoring rules at columns: {empty_rules}")

    formula_rows = 0
    detail_object_names: list[str] = []
    start_row = (rule_row or 4) + 1
    if score_header in detail_headers:
        score_column = detail_headers[score_header]
        for row in range(start_row, detail.max_row + 1):
            object_name = detail.cell(row, 1).value
            if not object_name:
                continue
            detail_object_names.append(str(object_name))
            for column in range(2, score_column):
                score = detail.cell(row, column).value
                if score is None:
                    if not contract or not any(item["name"] == object_name and item.get("status") in {"暂不评分", "不可比较"} and item.get("missing_reasons", {}).get(contract["indicators"][column - 2]["id"]) for item in contract["objects"]):
                        errors.append(f"Missing indicator score at 指标详表 row {row}, column {column} without documented unscored status")
                elif not is_formula(score) and (not isinstance(score, (int, float)) or not 0 <= float(score) <= 100):
                    errors.append(f"Invalid indicator score at 指标详表 row {row}, column {column}")
            composite = detail.cell(row, score_column).value
            if is_formula(composite):
                formula_rows += 1
                if contract:
                    expected = score_formula(contract["indicators"], row)
                    if normalized_formula(composite) != normalized_formula(expected):
                        errors.append(f"Composite formula must reference editable weights at 指标详表 row {row}")
            else:
                errors.append(f"Missing composite-score formula at 指标详表 row {row}")
    if formula_rows == 0:
        errors.append("No object score formulas found")

    overview = wb["模型总览"]
    overview_header_row = find_row(overview, "排名")
    overview_comparison_row = find_row(overview, "锚点定位")
    if not overview_comparison_row or overview.cell(overview_comparison_row, 4).value != "综合分差（百分制）":
        errors.append("模型总览 lacks percentage-scale anchor difference")
    overview_headers = header_map(overview, overview_header_row) if overview_header_row else {}
    required_overview = ("排名", "对象", "综合得分（百分制）", "评级", "模型置信度", "模型状态")
    if not overview_header_row or any(key not in overview_headers for key in required_overview):
        errors.append("模型总览 lacks complete result table")
    else:
        first = overview_header_row + 1
        last = overview_header_row + len(detail_object_names)
        score_letter = get_column_letter(overview_headers["综合得分（百分制）"])
        for row in range(first, last + 1):
            rank = overview.cell(row, overview_headers["排名"]).value
            object_name = overview.cell(row, overview_headers["对象"]).value
            expected_rank = f'=IF({score_letter}{row}="","",RANK({score_letter}{row},${score_letter}${first}:${score_letter}${last},0))'
            if normalized_formula(rank) != normalized_formula(expected_rank):
                errors.append(f"Wrong rank reference at 模型总览 row {row}")
            score = overview.cell(row, overview_headers["综合得分（百分制）"]).value
            rating = overview.cell(row, overview_headers["评级"]).value
            if contract and row - first < len(contract["objects"]):
                obj = contract["objects"][row - first]
                detail_row = start_row + row - first
                detail_score_letter = get_column_letter(score_column)
                expected_score = f'=IF(\'指标详表\'!{detail_score_letter}{detail_row}="","",\'指标详表\'!{detail_score_letter}{detail_row})'
                if object_name != obj["name"] or normalized_formula(score) != normalized_formula(expected_score):
                    errors.append(f"Wrong object or score reference at 模型总览 row {row}")
                for col, indicator in enumerate(contract["indicators"], start=3):
                    ref = f"'指标详表'!{get_column_letter(col - 1)}{detail_row}"
                    if normalized_formula(ref) not in normalized_formula(overview.cell(row, col).value):
                        errors.append(f"Wrong indicator reference at 模型总览 row {row}, column {col}")
                expected_rating = f'=IF({score_letter}{row}="","",{rating_formula(f"{score_letter}{row}", contract["model"].get("rating_rules", []))[1:]})'
                if rating != expected_rating:
                    errors.append(f"Rating rule mismatch at 模型总览 row {row}")
                if overview.cell(row, overview_headers["模型置信度"]).value != coverage_label(contract, "对象", obj["id"]):
                    errors.append(f"Confidence/coverage mismatch at 模型总览 row {row}")
                if overview.cell(row, overview_headers["模型状态"]).value != obj.get("status", contract["model"].get("status", "")):
                    errors.append(f"Model status mismatch at 模型总览 row {row}")
            elif not is_formula(score):
                errors.append(f"Missing percentage-score formula at 模型总览 row {row}")
            if not is_formula(rating) or f'IF({score_letter}{row}="","",' not in rating:
                errors.append(f"Missing or unguarded rating formula at 模型总览 row {row}")
            if rating == "待重算后判定":
                errors.append(f"Placeholder rating remains at 模型总览 row {row}")

    anchors = wb["锚点参照与赋分意义"]
    for required_label in ("锚点说明", "目标—锚点差距", "关键解读"):
        if not find_row(anchors, required_label):
            errors.append(f"锚点参照与赋分意义 lacks section: {required_label}")
    anchor_header_row = find_row(anchors, "维度")
    if not anchor_header_row:
        errors.append("锚点参照与赋分意义 lacks anchor score table")
    difference_header_row = find_row(anchors, "目标—锚点差距")
    anchor_total_row = find_row(anchors, "综合得分（百分制）")
    if not anchor_total_row:
        errors.append("锚点参照与赋分意义 lacks percentage-scale composite row")
    if not difference_header_row or anchors.cell(difference_header_row, 3 + (len(contract["indicators"]) if contract else max(len(detail_headers) - 2, 0))).value != "综合分差（百分制）":
        errors.append("锚点参照与赋分意义 lacks percentage-scale composite difference")

    sources = wb["数据溯源"]
    source_headers = header_map(sources, 2)
    required_sources = ("编号", "主体类型", "主体", "指标", "数据/主张", "数值", "单位", "数据状态", "URL", "备注")
    for required in required_sources:
        if required not in source_headers:
            errors.append(f"数据溯源 lacks column: {required}")
    source_mappings: set[tuple[str, str, str]] = set()
    source_rows: dict[tuple[str, str, str], list[int]] = {}
    if all(key in source_headers for key in required_sources):
        for row in range(3, sources.max_row + 1):
            claim = sources.cell(row, source_headers["数据/主张"]).value
            subject_type = sources.cell(row, source_headers["主体类型"]).value
            subject_name = sources.cell(row, source_headers["主体"]).value
            indicator = sources.cell(row, source_headers["指标"]).value
            url = sources.cell(row, source_headers["URL"]).value
            notes = sources.cell(row, source_headers["备注"]).value
            state = sources.cell(row, source_headers["数据状态"]).value
            if not claim or not state or not (url or notes):
                errors.append(f"Source row {row} lacks claim, status, or URL/source explanation")
            if subject_type and subject_name and indicator:
                pair = (str(subject_type), str(subject_name), str(indicator))
                source_mappings.add(pair)
                source_rows.setdefault(pair, []).append(row)

    assumptions = wb["模型假设"]
    for field in ("比较宇宙", "模型状态", "版本"):
        row = find_row(assumptions, field)
        if not row or not assumptions.cell(row, 2).value:
            errors.append(f"模型假设 lacks value for {field}")

    if contract:
        indicator_names = [item["name"] for item in contract.get("indicators", [])]
        object_names = [item["name"] for item in contract.get("objects", [])]
        anchor_names = [item["name"] for item in contract.get("anchors", [])]
        missing_indicators = [name for name in indicator_names if name not in detail_headers or name not in overview_headers]
        if missing_indicators:
            errors.append("Workbook lacks contract indicators: " + ", ".join(missing_indicators))
        if detail_object_names != object_names:
            errors.append("指标详表 objects do not match contract order")
        if rule_row:
            for col, indicator in enumerate(contract["indicators"], start=2):
                if detail.cell(rule_row, col).value != indicator.get("score_rule", ""):
                    errors.append(f"Scoring rule mismatch for {indicator['name']}")
                if detail.cell(weight_row, col).value != indicator["weight"]:
                    errors.append(f"Weight mismatch for {indicator['name']}")
                for i, obj in enumerate(contract["objects"], start=start_row):
                    if detail.cell(i, col).value != obj["scores"][indicator["id"]]:
                        errors.append(f"Score differs from contract for {obj['name']} / {indicator['name']}")
        if anchor_header_row:
            first_anchor_col = 6 + len(contract["objects"])
            actual_anchor_names = [cell.value for cell in anchors[anchor_header_row][first_anchor_col - 1:] if cell.value]
            missing_anchors = [name for name in anchor_names if name not in actual_anchor_names]
            if missing_anchors:
                errors.append("锚点参照与赋分意义 lacks contract anchors: " + ", ".join(missing_anchors))
            groups = dimension_groups(contract)
            dimension_by_id = {dim["id"]: (dim, weight) for dim, _, weight in groups}
            for row, indicator in enumerate(contract["indicators"], start=anchor_header_row + 1):
                dim, weight = dimension_by_id[indicator["dimension_id"]]
                members = [item for item in contract["indicators"] if item["dimension_id"] == indicator["dimension_id"]]
                expected = [dim["name"], dimension_weight_formula(contract, members), indicator["name"],
                            f'=IF(B{row}>0,E{row}/B{row},"")', "=" + indicator_weight_refs(contract, [indicator])[0]]
                for col, value in enumerate(expected, start=1):
                    actual = anchors.cell(row, col).value
                    equal = abs(actual - value) < 1e-10 if isinstance(value, (int, float)) and isinstance(actual, (int, float)) else normalized_formula(actual) == normalized_formula(value)
                    if not equal:
                        errors.append(f"Dimension hierarchy/weight mismatch at row {row}, column {col}")
                for col, obj in enumerate(contract["objects"], start=6):
                    ref = f"'指标详表'!{get_column_letter(row - anchor_header_row + 1)}{start_row + col - 6}"
                    expected_ref = f'=IF({ref}="","",{ref})'
                    if normalized_formula(anchors.cell(row, col).value) != normalized_formula(expected_ref):
                        errors.append(f"Wrong target reference in hierarchy for {obj['name']} / {indicator['name']}")
            for col, anchor in enumerate(contract["anchors"], start=first_anchor_col):
                for i, indicator in enumerate(contract["indicators"], start=anchor_header_row + 1):
                    if anchors.cell(i, col).value != anchor["scores"][indicator["id"]]:
                        errors.append(f"Anchor score differs from contract for {anchor['name']} / {indicator['name']}")
                if anchor_total_row:
                    letter = get_column_letter(col)
                    refs = [f"{letter}{anchor_header_row + 1 + i}" for i in range(len(contract["indicators"]))]
                    expected = weighted_formula(refs, indicator_weight_refs(contract))
                    if normalized_formula(anchors.cell(anchor_total_row, col).value) != normalized_formula(expected):
                        errors.append(f"Wrong percentage-scale anchor composite for {anchor['name']}")
            summary_row = find_row(anchors, "维度汇总（百分制）")
            overview_dimension_row = find_row(overview, "分维度比较（百分制）")
            if not summary_row or not overview_dimension_row:
                errors.append("Workbook lacks dimension results before overall results")
            else:
                subjects = contract["objects"] + contract["anchors"]
                for offset, (dimension, members, weight) in enumerate(groups, start=1):
                    r = summary_row + offset
                    overview_r = overview_dimension_row + offset
                    if anchors.cell(r, 1).value != dimension["name"] or overview.cell(overview_r, 1).value != dimension["name"]:
                        errors.append(f"Missing or misplaced dimension: {dimension['name']}")
                    for sheet, dimension_row in ((anchors, r), (overview, overview_r)):
                        actual_weight = sheet.cell(dimension_row, 2).value
                        if normalized_formula(actual_weight) != normalized_formula(dimension_weight_formula(contract, members)):
                            errors.append(f"Wrong dimension weight: {dimension['name']}")
                    for col, subject in enumerate(subjects, start=3):
                        refs = [f"{get_column_letter(col + 3)}{anchor_header_row + 1 + contract['indicators'].index(item)}" for item in members]
                        expected = weighted_formula(refs, indicator_weight_refs(contract, members))
                        ref = f"'锚点参照与赋分意义'!{get_column_letter(col)}{r}"
                        overview_expected = f'=IF({ref}="","",{ref})'
                        if normalized_formula(anchors.cell(r, col).value) != normalized_formula(expected):
                            errors.append(f"Wrong dimension result: {dimension['name']} / {subject['name']}")
                        if normalized_formula(overview.cell(overview_r, col).value) != normalized_formula(overview_expected):
                            errors.append(f"Wrong overview dimension reference: {dimension['name']} / {subject['name']}")
        expected_mappings = {
            ("对象", object_name, indicator_name)
            for object_name in object_names
            for indicator_name in indicator_names
        } | {
            ("锚点", anchor_name, indicator_name)
            for anchor_name in anchor_names
            for indicator_name in indicator_names
        }
        missing_mappings = expected_mappings - source_mappings
        if missing_mappings:
            errors.append(f"数据溯源 lacks {len(missing_mappings)} required subject-indicator mappings")
        source_by_id = {item["id"]: item for item in contract["sources"]}
        seen_source_ids: set[str] = set()
        for r in range(3, sources.max_row + 1):
            source_id = sources.cell(r, source_headers["编号"]).value
            if source_id not in source_by_id or source_id in seen_source_ids:
                errors.append(f"Unknown or duplicate source at 数据溯源 row {r}")
                continue
            seen_source_ids.add(source_id)
            item = source_by_id[source_id]
            subject_name = next(subject["name"] for subject in contract["objects" if item["subject_type"] == "对象" else "anchors"] if subject["id"] == item["subject_id"])
            indicator_name = next(indicator["name"] for indicator in contract["indicators"] if indicator["id"] == item["indicator_id"])
            expected = {"主体类型": item["subject_type"], "主体": subject_name, "指标": indicator_name, "数据/主张": item.get("claim", ""), "数值": item.get("value", ""), "单位": item.get("unit", ""), "时间": item.get("date", ""), "数据状态": item.get("status", ""), "审核评级": item.get("rating", ""), "URL": item.get("url", ""), "备注": item.get("notes", "")}
            for field, value in expected.items():
                if field in source_headers and (sources.cell(r, source_headers[field]).value or "") != (value or ""):
                    errors.append(f"Source {source_id} field {field} differs from contract")
        if seen_source_ids != set(source_by_id):
            errors.append("数据溯源 does not contain every contract source")
        for subject_type, items in (("对象", contract["objects"]), ("锚点", contract["anchors"])):
            for item in items:
                for indicator in contract["indicators"]:
                    pair = (subject_type, item["name"], indicator["name"])
                    score = item["scores"][indicator["id"]]
                    rows = source_rows.get(pair, [])
                    if score is None and not any(sources.cell(r, source_headers["数据状态"]).value == "数据缺口" for r in rows):
                        errors.append(f"No gap source for {item['name']} / {indicator['name']}")
        anchor_columns = {item["id"]: i for i, item in enumerate(contract["anchors"], start=6 + len(contract["objects"]))}
        object_rows = {item["id"]: start_row + i for i, item in enumerate(contract["objects"])}
        if difference_header_row and anchor_header_row:
            for offset, comparison in enumerate(contract.get("anchor_comparisons", []), start=1):
                r = difference_header_row + offset
                detail_row = object_rows[comparison["object_id"]]
                anchor_col = get_column_letter(anchor_columns[comparison["anchor_id"]])
                for i, indicator in enumerate(contract["indicators"]):
                    detail_col = get_column_letter(i + 2)
                    expected = difference_formula(f"'指标详表'!{detail_col}{detail_row}", f"{anchor_col}{anchor_header_row + 1 + i}")
                    if normalized_formula(anchors.cell(r, i + 3).value) != normalized_formula(expected):
                        errors.append(f"Wrong anchor difference formula at row {r}, indicator {indicator['name']}")
                if anchor_total_row:
                    detail_score_col = get_column_letter(score_column)
                    expected = difference_formula(f"'指标详表'!{detail_score_col}{detail_row}", f"{anchor_col}{anchor_total_row}")
                    if normalized_formula(anchors.cell(r, len(contract["indicators"]) + 3).value) != normalized_formula(expected):
                        errors.append(f"Wrong percentage-scale anchor difference at row {r}")
                    if overview_comparison_row:
                        overview_r = overview_comparison_row + offset
                        overview_expected = difference_formula(f"'指标详表'!{detail_score_col}{detail_row}", f"'锚点参照与赋分意义'!{anchor_col}{anchor_total_row}")
                        if normalized_formula(overview.cell(overview_r, 4).value) != normalized_formula(overview_expected):
                            errors.append(f"Wrong percentage-scale overview anchor difference at row {overview_r}")
        reasons_row = find_row(assumptions, "对象")
        if not reasons_row:
            errors.append("模型假设 lacks per-object scoring reasons")
        else:
            r = reasons_row + 1
            for subject in contract["objects"] + contract["anchors"]:
                for indicator in contract["indicators"]:
                    expected_reason = (subject.get("missing_reasons", {}).get(indicator["id"], "") if subject["scores"][indicator["id"]] is None
                                       else subject.get("score_reasons", {}).get(indicator["id"], ""))
                    if (assumptions.cell(r, 1).value != subject["name"] or
                            assumptions.cell(r, 3).value != indicator["name"] or
                            assumptions.cell(r, 5).value != expected_reason):
                        errors.append(f"Missing or mismatched scoring reason: {subject['name']} / {indicator['name']}")
                    r += 1
        if contract["model"].get("mode") == "组合加权":
            if "组合加权" not in wb.sheetnames:
                errors.append("组合加权 mode lacks combination sheet")
            else:
                combo = wb["组合加权"]
                if combo.cell(2, len(contract["indicators"]) + 4).value != "综合得分（百分制）":
                    errors.append("组合加权 lacks percentage-scale composite column")
                total_row = len(contract["objects"]) + 3
                component_names = {item["id"]: item["name"] for item in contract["objects"]}
                for r, component in enumerate(contract["combination"]["components"], start=3):
                    actual_weight = combo.cell(r, 2).value
                    if (combo.cell(r, 1).value != component_names[component["object_id"]]
                            or not isinstance(actual_weight, (int, float))
                            or abs(actual_weight - float(component["weight"])) > 1e-12
                            or combo.cell(r, 3).value != component["source"]):
                        errors.append(f"Combination component differs from contract at row {r}")
                    detail_row = object_rows[component["object_id"]]
                    detail_score_col = get_column_letter(score_column)
                    expected_score = f"=指标详表!{detail_score_col}{detail_row}"
                    if normalized_formula(combo.cell(r, len(contract["indicators"]) + 4).value) != expected_score:
                        errors.append(f"Wrong percentage-scale combination composite at row {r}")
                if combo.cell(total_row, 2).value != f"=SUM(B3:B{total_row - 1})":
                    errors.append("Combination weight total formula is wrong")
                for col in range(4, len(contract["indicators"]) + 5):
                    letter = get_column_letter(col)
                    if combo.cell(total_row, col).value != f"=SUMPRODUCT($B$3:$B${total_row - 1},{letter}3:{letter}{total_row - 1})":
                        errors.append(f"Combination weighted formula is wrong at column {letter}")

    cached = load_workbook(path, data_only=True)
    for sheet in cached:
        for row in sheet:
            for cell in row:
                if cell.data_type == "e":
                    errors.append(f"Cached Excel error at {sheet.title}!{cell.coordinate}: {cell.value}")
    warnings.append("Formula structure and available cached results checked; this validator does not recalculate formulas")
    return {
        "passed": not errors,
        "errors": errors,
        "warnings": warnings,
        "sheets": wb.sheetnames,
        "formula_rows": formula_rows,
        "source_mappings": len(source_mappings),
    }


def main() -> int:
    parser = argparse.ArgumentParser()
    parser.add_argument("--workbook", required=True, type=Path)
    parser.add_argument("--contract", type=Path, help="Optional JSON input contract for cross-checking")
    args = parser.parse_args()
    result = validate(args.workbook, args.contract)
    print(json.dumps(result, ensure_ascii=False, indent=2))
    return 0 if result["passed"] else 1


if __name__ == "__main__":
    raise SystemExit(main())
